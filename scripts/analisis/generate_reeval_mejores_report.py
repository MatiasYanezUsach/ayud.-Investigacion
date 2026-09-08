#!/usr/bin/env python3
"""Consolidate the re-evaluation of the best algorithm of each hybrid condition.

Input (written by scripts/windows/run_reeval_mejores.bat or scripts/linux/run_reeval_mejores.sh):
  out/reeval_mejores/<Bxx>/evolution0/MISPResults.out            wall-clock ms, cost, rel. error, hits per instance
  out/reeval_mejores/<Bxx>/evolution0/job.0.CplexUsage.summary.csv  one row per instance: TotalCalls, TotalTimeUsed (CPLEX CPU s)
  out/reeval_mejores/<Bxx>/evolution0/job.0.CplexUsage.detailed.csv one row per CPLEX call (cross-check)
  out/baseline/cplex_baseline_per_instance.csv                    T_base per instance (wall-clock s)

Output: reportes/REEVALUACION_MEJORES_ALGORITMOS.xlsx

Usage: python scripts/analisis/generate_reeval_mejores_report.py [--base DIR] [--out FILE]
"""
import os
# Todas las rutas del script son relativas a la raiz del repositorio (dos niveles arriba de scripts/analisis).
os.chdir(os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")))

import argparse
import csv
import glob
import re
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument("--base", default="out/reeval_mejores", help="carpeta con las corridas Bxx/evolution0")
ap.add_argument("--out", default="reportes/REEVALUACION_MEJORES_ALGORITMOS.xlsx")
args = ap.parse_args()

LABELS = OrderedDict([("B10", 0.10), ("B25", 0.25), ("B50", 0.50), ("B75", 0.75), ("B100", 1.00)])
INSTANCES = ["3C_20_50-02.txt", "3C_20_50-03.txt", "3C_20_66-01.txt", "3C_20_66-02.txt",
             "3C_20_66-03.txt", "3C_20_80-01.txt", "3C_20_80-02.txt", "3C_20_80-03.txt"]

# Valores del experimento original (evalthreads = 6), Cap. 4 y TIEMPO_CPLEX_MEJOR_ALGORITMO_POR_GRUPO.xlsx
ORIGINAL = {
    "B10":  dict(erp=0.0069106, hits=3, cplex=13.250,  wall=14.215,  origen="grupo1 ejec 2 gen 16"),
    "B25":  dict(erp=0.0048098, hits=4, cplex=80.095,  wall=41.015,  origen="grupo2 ejec 1 gen 48"),
    "B50":  dict(erp=0.0031093, hits=5, cplex=90.172,  wall=105.482, origen="grupo3 ejec 4 gen 53"),
    "B75":  dict(erp=0.0010745, hits=6, cplex=228.626, wall=169.213, origen="grupo4 ejec 2 gen 9"),
    "B100": dict(erp=0.0,       hits=8, cplex=404.328, wall=174.693, origen="grupo5 ejec 2 gen 49"),
}

# MISPResults.out record (PDPProblemEvo): gen individual wall_ms cost optimal rel_err depth size ideal_nodes ERL fitness hits
N = r"-?\d+(?:\.\d+)?(?:E-?\d+)?"
PAT = re.compile(r"(\d+) (ec\.gp\.GPIndividual@\d+\{-?\d+\}) (%s) (%s) (%s) (%s) (\d+) (\d+) (%s) (%s) (%s) (\d+)" % ((N,) * 7))


def baseline():
    b = {}
    with open("out/baseline/cplex_baseline_per_instance.csv") as f:
        next(f)
        for line in f:
            p = line.strip().split(",")
            if len(p) >= 2:
                b[p[0]] = float(p[1])
    return b


def first_existing(*paths):
    for p in paths:
        if os.path.exists(p):
            return p
    return None


def read_run(label):
    d = f"{args.base}/{label}/evolution0"
    res = first_existing(f"{d}/MISPResults.out", f"{d}/job.0.MISPResults.out")
    if res is None:
        return None
    txt = open(res, encoding="utf-8", errors="replace").read()
    recs = [m.groups() for m in PAT.finditer(txt)]
    if not recs:
        return None
    inds = sorted({r[1] for r in recs})
    if len(inds) != 1:
        print(f"ADVERTENCIA {label}: se esperaba 1 individuo y hay {len(inds)}; se usa el primero")
    ind = inds[0]
    rows = [r for r in recs if r[1] == ind]
    by_inst = {}
    # PDPProblemEvo evaluates instances in the order of data/evolution (the first 8): map by position
    for k, r in enumerate(rows[:len(INSTANCES)]):
        by_inst[INSTANCES[k]] = dict(wall=float(r[2]) / 1000.0, cost=float(r[3]), opt=float(r[4]), rel=float(r[5]),
                                    depth=int(r[6]), size=int(r[7]), hits=int(r[11]))
    # CPLEX per instance from summary.csv (one row per instance for a single individual)
    cplex_t, calls = defaultdict(float), defaultdict(int)
    summ = f"{d}/job.0.CplexUsage.summary.csv"
    if os.path.exists(summ):
        with open(summ) as f:
            for row in csv.DictReader(f):
                if row["Individual"] != ind:
                    continue
                cplex_t[row["Instance"]] += float(row["TotalTimeUsed"])
                calls[row["Instance"]] += int(row["TotalCalls"])
    # cross-check with detailed.csv
    det_t, det_n = defaultdict(float), defaultdict(int)
    det = f"{d}/job.0.CplexUsage.detailed.csv"
    if os.path.exists(det):
        with open(det) as f:
            for row in csv.DictReader(f):
                if row["Individual"] != ind:
                    continue
                det_t[row["Instance"]] += float(row["TimeUsed"])
                det_n[row["Instance"]] += 1
    if not cplex_t and det_t:
        cplex_t, calls = det_t, det_n
    total_line = txt.strip().splitlines()[-1].strip()
    run_total = float(total_line) if re.fullmatch(N, total_line) else None
    return dict(ind=ind, by_inst=by_inst, cplex=cplex_t, calls=calls, det=det_t, det_n=det_n, run_total_ms=run_total)


bl = baseline()
runs = OrderedDict((lab, read_run(lab)) for lab in LABELS)
found = [lab for lab, r in runs.items() if r]
if not found:
    raise SystemExit(f"No se encontraron corridas en {args.base}/<Bxx>/evolution0/. Ejecute primero run_reeval_mejores.")
missing = [lab for lab, r in runs.items() if not r]
if missing:
    print("ADVERTENCIA: sin datos para", ", ".join(missing))

# ---------- Excel ----------
wb = openpyxl.Workbook()
bold = Font(bold=True)
hdr_fill = PatternFill("solid", fgColor="DDEBF7")
tot_fill = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, r, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(r, c)
        cell.font = bold; cell.fill = hdr_fill; cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autowidth(ws):
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, w + 2), 60)


def matrix_sheet(ws, title, note, value_of, fmt="0.000"):
    ws.cell(1, 1).value = note
    ws.cell(1, 1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(found))
    ws.row_dimensions[1].height = 75
    hdr = ["Instancia", "T_base (s)"] + found
    for c, h in enumerate(hdr, 1):
        ws.cell(3, c).value = h
    style_header(ws, 3, len(hdr))
    r0 = 4
    col_tot = defaultdict(float)
    for i, inst in enumerate(INSTANCES):
        ws.cell(r0 + i, 1).value = inst.replace(".txt", "")
        ws.cell(r0 + i, 2).value = round(bl.get(inst, 0.0), 3)
        col_tot[2] += bl.get(inst, 0.0)
        for j, lab in enumerate(found):
            v = value_of(runs[lab], inst)
            ws.cell(r0 + i, 3 + j).value = round(v, 3)
            col_tot[3 + j] += v
    rt = r0 + len(INSTANCES)
    ws.cell(rt, 1).value = "Total (8 instancias)"
    for c in range(2, 3 + len(found)):
        ws.cell(rt, c).value = round(col_tot[c], 3)
    ws.cell(rt + 1, 1).value = "Ganancia vs T_base (%)"
    for j in range(len(found)):
        ws.cell(rt + 1, 3 + j).value = (col_tot[2] - col_tot[3 + j]) / col_tot[2] if col_tot[2] else None
        ws.cell(rt + 1, 3 + j).number_format = "0.0%"
    ws.cell(rt + 2, 1).value = f"TOTAL DE LOS {len(found)} ALGORITMOS (s)"
    ws.cell(rt + 2, 3).value = round(sum(col_tot[3 + j] for j in range(len(found))), 3)
    ws.merge_cells(start_row=rt + 2, start_column=3, end_row=rt + 2, end_column=2 + len(found))
    for row in ws.iter_rows(min_row=3, max_row=rt + 2, max_col=len(hdr)):
        for cell in row:
            cell.border = border
            if cell.column >= 2 and cell.row <= rt:
                cell.number_format = fmt
    for c in range(1, len(hdr) + 1):
        for r in (rt, rt + 2):
            ws.cell(r, c).font = bold; ws.cell(r, c).fill = tot_fill
    autowidth(ws)
    return col_tot


ws1 = wb.active
ws1.title = "Tiempo CPLEX por instancia"
tot_cplex = matrix_sheet(ws1, "cplex",
    "Tiempo de CPLEX (s, CPU del solver, ClockType=1, Threads=2) consumido por el mejor algoritmo de cada condicion al "
    "resolver cada una de las 8 instancias, medido en la re-evaluacion con un solo hilo (evalthreads=1) y una JVM por "
    "algoritmo, cada uno con su propio presupuesto. Fuente: job.0.CplexUsage.summary.csv (TotalTimeUsed). "
    "T_base = CPLEX puro por instancia (out/baseline/cplex_baseline_per_instance.csv).",
    lambda r, inst: r["cplex"].get(inst, 0.0))

ws2 = wb.create_sheet("Tiempo total por instancia")
tot_wall = matrix_sheet(ws2, "wall",
    "Tiempo TOTAL de ejecucion del mejor algoritmo de cada condicion sobre cada instancia (s, reloj de pared): desde que "
    "el arbol empieza a evaluarse hasta que termina, incluidas heuristicas y llamadas a CPLEX. Fuente: MISPResults.out, "
    "columna 3 (ms, System.nanoTime en PDPProblemEvo). Sin contencion: un individuo, un hilo de evaluacion.",
    lambda r, inst: r["by_inst"].get(inst, {}).get("wall", 0.0))

ws3 = wb.create_sheet("Resumen por algoritmo")
hdr3 = ["Condicion", "Origen (experimento)", "Presupuesto", "ID individuo (ECJ, re-eval)", "Tamano arbol",
        "ERP re-eval", "ERP original", "Hits re-eval (de 8)", "Hits original", "Llamadas CPLEX (8 inst.)",
        "Tiempo CPLEX re-eval (s)", "Tiempo CPLEX original (s)", "Tiempo total re-eval (s, pared)",
        "Tiempo total original (s, pared)", "Duracion corrida completa (s)"]
for c, h in enumerate(hdr3, 1):
    ws3.cell(1, c).value = h
style_header(ws3, 1, len(hdr3))
r = 2
sum_c, sum_w = 0.0, 0.0
for lab in found:
    run = runs[lab]; o = ORIGINAL[lab]
    insts = [run["by_inst"][i] for i in INSTANCES if i in run["by_inst"]]
    erp = sum(x["rel"] for x in insts) / len(insts) if insts else None
    hits = sum(1 for x in insts if x["rel"] == 0.0)
    c_tot = sum(run["cplex"].get(i, 0.0) for i in INSTANCES)
    w_tot = sum(x["wall"] for x in insts)
    sum_c += c_tot; sum_w += w_tot
    vals = [lab, o["origen"], LABELS[lab], run["ind"], insts[0]["size"] if insts else None, erp, o["erp"], hits, o["hits"],
            sum(run["calls"].get(i, 0) for i in INSTANCES), round(c_tot, 3), o["cplex"], round(w_tot, 3), o["wall"],
            round(run["run_total_ms"] / 1000.0, 3) if run["run_total_ms"] else None]
    for c, v in enumerate(vals, 1):
        ws3.cell(r, c).value = v; ws3.cell(r, c).border = border
    ws3.cell(r, 3).number_format = "0%"; ws3.cell(r, 6).number_format = "0.000000"; ws3.cell(r, 7).number_format = "0.000000"
    r += 1
ws3.cell(r, 1).value = f"TOTAL {len(found)} algoritmos"; ws3.cell(r, 1).font = bold
ws3.cell(r, 11).value = round(sum_c, 3); ws3.cell(r, 13).value = round(sum_w, 3)
ws3.cell(r, 12).value = round(sum(ORIGINAL[l]["cplex"] for l in found), 3); ws3.cell(r, 14).value = round(sum(ORIGINAL[l]["wall"] for l in found), 3)
for c in range(1, len(hdr3) + 1):
    ws3.cell(r, c).font = bold; ws3.cell(r, c).fill = tot_fill; ws3.cell(r, c).border = border
ws3.cell(r + 2, 1).value = ("Nota: 'original' = valores del experimento (evalthreads = 6, logger compartido) reportados en el Cap. 4 y en "
                            "TIEMPO_CPLEX_MEJOR_ALGORITMO_POR_GRUPO.xlsx. Si el ERP o los hits de la re-evaluacion difieren del original, "
                            "la causa esperable es el presupuesto efectivo: en el experimento seis hilos compartian el presupuesto de CPLEX.")
autowidth(ws3)

ws4 = wb.create_sheet("Detalle por instancia")
hdr4 = ["Condicion", "Instancia", "Optimo", "Costo obtenido", "Error relativo", "Llamadas CPLEX", "Tiempo CPLEX (s, summary)",
        "Tiempo CPLEX (s, detailed)", "Tiempo total (s, pared)", "T_base (s)", "Presupuesto asignado (s)"]
for c, h in enumerate(hdr4, 1):
    ws4.cell(1, c).value = h
style_header(ws4, 1, len(hdr4))
r = 2
for lab in found:
    run = runs[lab]
    for inst in INSTANCES:
        x = run["by_inst"].get(inst)
        if not x:
            continue
        vals = [lab, inst.replace(".txt", ""), x["opt"], x["cost"], x["rel"], run["calls"].get(inst, 0),
                round(run["cplex"].get(inst, 0.0), 3), round(run["det"].get(inst, 0.0), 3), round(x["wall"], 3),
                round(bl.get(inst, 0.0), 3), round(bl.get(inst, 0.0) * LABELS[lab], 3)]
        for c, v in enumerate(vals, 1):
            ws4.cell(r, c).value = v; ws4.cell(r, c).border = border
        ws4.cell(r, 5).number_format = "0.000000"
        r += 1
autowidth(ws4)

ws0 = wb.create_sheet("Leeme", 0)
ws0.column_dimensions["A"].width = 32; ws0.column_dimensions["B"].width = 110
rows0 = [
    ("Que es", "Re-evaluacion controlada de los 5 mejores algoritmos generados (uno por condicion hibrida B10..B100) sobre las 8 "
               "instancias del protocolo, pedida por el profesor guia (2026-09-08). Cada algoritmo se carga como poblacion fija en ECJ "
               "(pop.file, generations=1) y se evalua con su propio presupuesto de CPLEX."),
    ("Por que se repite la medicion", "En el experimento original se evaluaban 6 individuos en paralelo y CplexTerminal/CplexUsageLogger "
               "guardan estado estatico compartido (limitacion del Cap. 5). Aqui: un individuo, un hilo, una JVM por algoritmo. "
               "Los tiempos quedan bien atribuidos y sin contencion."),
    ("Tiempo de CPLEX", "CPU del solver (ClockType=1, Threads=2) sumado sobre todas las llamadas del arbol en la instancia. Hoja 'Tiempo CPLEX por instancia'."),
    ("Tiempo total", "Reloj de pared de toda la evaluacion del arbol en la instancia (heuristicas + CPLEX). Hoja 'Tiempo total por instancia'. "
               "Comparable con T_base, que tambien es reloj de pared."),
    ("Total de los 5", "Fila 'TOTAL DE LOS 5 ALGORITMOS' en ambas hojas: suma de los totales por algoritmo sobre las 8 instancias."),
    ("Comparacion con el experimento", "Hoja 'Resumen por algoritmo': columnas 'original' con los valores publicados en el Cap. 4 "
               "(TIEMPO_CPLEX_MEJOR_ALGORITMO_POR_GRUPO.xlsx)."),
    ("Como se genero", "scripts/windows/run_reeval_mejores.bat (o scripts/linux/run_reeval_mejores.sh) y luego "
               "scripts/analisis/generate_reeval_mejores_report.py. Poblacion en out/reeval_mejores/poblacion/. Salida cruda en out/reeval_mejores/<Bxx>/evolution0/."),
    ("Semilla", "seed.0 = 20260908 (fija; no afecta la evaluacion de un individuo ya construido salvo por RandomMove, si aparece en el arbol)."),
]
ws0.cell(1, 1).value = "Re-evaluacion de los mejores algoritmos por condicion"; ws0.cell(1, 1).font = Font(bold=True, size=14)
for i, (k, v) in enumerate(rows0, 3):
    ws0.cell(i, 1).value = k; ws0.cell(i, 1).font = bold; ws0.cell(i, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws0.cell(i, 2).value = v; ws0.cell(i, 2).alignment = Alignment(vertical="top", wrap_text=True)
    ws0.row_dimensions[i].height = max(30, 15 * (len(v) // 105 + 1))

os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
wb.save(args.out)
print("OK ->", args.out)
for lab in found:
    print(f"  {lab}: CPLEX={sum(runs[lab]['cplex'].get(i,0) for i in INSTANCES):.3f} s  pared={sum(x['wall'] for x in runs[lab]['by_inst'].values()):.3f} s")
print(f"  TOTAL {len(found)}: CPLEX={sum_c:.3f} s  pared={sum_w:.3f} s")
