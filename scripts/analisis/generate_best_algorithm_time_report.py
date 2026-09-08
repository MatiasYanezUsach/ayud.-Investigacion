#!/usr/bin/env python3
"""Build the per-instance time report for the best generated algorithm of each condition.

Sources (raw experiment output, no manual edits):
  out/results/grupoN/evolutionM/job.M.EstadisticaProm&Mej.csv  -> best individual (min BestERP over all gens/runs)
  out/results/grupoN/evolutionM/job.M.MISPResults.out          -> quality and wall-clock time per instance
  out/results/grupoN/evolutionM/job.M.CplexUsage.detailed.csv  -> CPLEX time per instance (CPU clock)
  out/baseline/cplex_baseline_results.csv                      -> pure-CPLEX baseline per instance

Output: TIEMPO_CPLEX_MEJOR_ALGORITMO_POR_GRUPO.xlsx (Spanish sheet names/labels, matching the other reports).
"""
import os
# Todas las rutas del script son relativas a la raiz del repositorio (dos niveles arriba de scripts/analisis).
os.chdir(os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")))
import csv
import glob
import os
import re
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = "out/results"
OUT = "reportes/TIEMPO_CPLEX_MEJOR_ALGORITMO_POR_GRUPO.xlsx"
GROUPS = OrderedDict([(1, "B10 (10%)"), (2, "B25 (25%)"), (3, "B50 (50%)"), (4, "B75 (75%)"), (5, "B100 (100%)")])
INSTANCES = ["3C_20_50-02.txt", "3C_20_50-03.txt", "3C_20_66-01.txt", "3C_20_66-02.txt",
             "3C_20_66-03.txt", "3C_20_80-01.txt", "3C_20_80-02.txt", "3C_20_80-03.txt"]

# MISPResults.out record: [gen] [individual] [wall ms] [profit] [optimal] [rel err] [depth] [size] [ideal nodes] [ERL] [fitness] [hits]
# Records may be concatenated on one line (concurrent writers), so parse with a strict regex over the whole file.
N = r"-?\d+(?:\.\d+)?(?:E-?\d+)?"
PAT = re.compile(r"(\d+) (ec\.gp\.GPIndividual@\d+\{-?\d+\}) (%s) (%s) (%s) (%s) (\d+) (\d+) (%s) (%s) (%s) (\d+)" % (N, N, N, N, N, N, N))


def pf(s):
    return float(s.replace(",", "."))


def baseline():
    b = {}
    with open("out/baseline/cplex_baseline_results.csv") as f:
        next(f)
        for line in f:
            p = line.strip().split(",")
            # Type,Instance,Optimal(int,dec),Cost(int,dec),TimeSeconds(int,dec),Status,Gap(int,dec)
            if len(p) < 9 or p[0] != "Evolution":
                continue
            b[p[1]] = float(p[6] + "." + p[7])
    return b


def best_of_group(g):
    best = None
    for ev in sorted(glob.glob(f"{BASE}/grupo{g}/evolution*")):
        m = int(os.path.basename(ev).replace("evolution", ""))
        with open(f"{ev}/job.{m}.EstadisticaProm&Mej.csv", encoding="latin-1") as f:
            for row in csv.DictReader(f, delimiter=";"):
                erp = pf(row["BestERP"])
                gen = int(row["Gen"])
                if best is None or erp < best[0]:
                    best = (erp, m, gen, int(row["BestHits"]), int(row["BestSize"]))
    return best


def best_individual(g, m, state_gen):
    acc = defaultdict(list)
    wall = defaultdict(list)
    profit = defaultdict(list)
    txt = open(f"{BASE}/grupo{g}/evolution{m}/job.{m}.MISPResults.out").read()
    for p in (mm.groups() for mm in PAT.finditer(txt)):
        if int(p[0]) != state_gen:
            continue
        acc[p[1]].append(float(p[5]))
        wall[p[1]].append(float(p[2]))
        profit[p[1]].append((float(p[3]), float(p[4])))
    ind = min(acc, key=lambda k: sum(acc[k]) / len(acc[k]))
    return ind, acc[ind], wall[ind], profit[ind]


def cplex_time(g, m, state_gen, ind):
    """t   = BudgetUsed (cumulative) of the last call sequence per instance (criterion of table tab:ganancia-tiempo).
    raw = sum of TimeUsed over ALL rows attributed to the individual (differs only when sequences are duplicated
          by concurrent attribution in the shared logger)."""
    t = defaultdict(float)
    raw = defaultdict(float)
    calls = defaultdict(int)
    seqs = defaultdict(int)
    with open(f"{BASE}/grupo{g}/evolution{m}/job.{m}.CplexUsage.detailed.csv") as f:
        for row in csv.DictReader(f):
            if int(row["Generation"]) == state_gen and row["Individual"] == ind:
                i = row["Instance"]
                if row["CallNumber"] == "1":
                    seqs[i] += 1
                    calls[i] = 0
                calls[i] += 1
                raw[i] += float(row["TimeUsed"])
                t[i] = float(row["BudgetUsed"])
    return t, calls, raw, seqs


bl = baseline()
rows = []
for g, label in GROUPS.items():
    erp, m, gen, hits, size = best_of_group(g)
    ind, rel, wall, prof = best_individual(g, m, gen - 1)
    t, calls, raw, seqs = cplex_time(g, m, gen - 1, ind)
    rows.append(dict(g=g, label=label, evo=m, ejec=m + 1, gen=gen, erp=erp, hits=hits, size=size, ind=ind,
                     t=t, calls=calls, raw=raw, seqs=seqs, rel=rel, wall=wall, prof=prof))
    print(f"grupo{g} {label}: evo{m} gen{gen} ERP={erp:.6f} total CPLEX={sum(t.values()):.3f}s")

# ---------- Excel ----------
wb = openpyxl.Workbook()
bold = Font(bold=True)
hdr_fill = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, r, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(r, c)
        cell.font = bold
        cell.fill = hdr_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autowidth(ws):
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, w + 2), 60)


def summary_sheet(ws, note, value_of, note_height):
    """8 instances x (baseline + 5 conditions) with totals and gains."""
    ws.cell(1, 1).value = note
    ws.cell(1, 1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(rows))
    ws.row_dimensions[1].height = note_height
    hdr = ["Instancia", "Baseline CPLEX (s)"] + [r["label"] for r in rows]
    for c, h in enumerate(hdr, 1):
        ws.cell(3, c).value = h
    style_header(ws, 3, len(hdr))
    r0 = 4
    for i, inst in enumerate(INSTANCES):
        ws.cell(r0 + i, 1).value = inst.replace(".txt", "")
        ws.cell(r0 + i, 2).value = round(bl[inst], 3)
        for j, r in enumerate(rows):
            ws.cell(r0 + i, 3 + j).value = value_of(r, inst, i)
    rt = r0 + len(INSTANCES)
    # Totals and gains are written as values (not formulas) so that readers without a
    # spreadsheet engine (pandas, openpyxl, LibreOffice headless) see them too.
    ws.cell(rt, 1).value = "Total"
    ws.cell(rt, 1).font = bold
    totals = {}
    for c in range(2, 3 + len(rows)):
        totals[c] = round(sum(ws.cell(r, c).value for r in range(r0, rt)), 3)
        ws.cell(rt, c).value = totals[c]
        ws.cell(rt, c).font = bold
    ws.cell(rt + 1, 1).value = "Ganancia vs baseline (s)"
    ws.cell(rt + 2, 1).value = "Ganancia (%)"
    for j in range(len(rows)):
        c = 3 + j
        ws.cell(rt + 1, c).value = round(totals[2] - totals[c], 3)
        ws.cell(rt + 2, c).value = (totals[2] - totals[c]) / totals[2]
        ws.cell(rt + 2, c).number_format = "0.0%"
    for row in ws.iter_rows(min_row=3, max_row=rt + 2, max_col=len(hdr)):
        for cell in row:
            cell.border = border
            if cell.column >= 2 and cell.row <= rt + 1:
                cell.number_format = "0.000"
    autowidth(ws)


# Sheet: tiempo CPLEX por instancia (thesis table criterion)
ws = wb.active
ws.title = "Tiempo CPLEX por instancia"
summary_sheet(ws,
              "Tiempo de CPLEX (s, CPU, Threads=2) consumido por el MEJOR algoritmo de cada grupo al evaluar cada una de las 8 instancias, "
              "en la generación en que fue encontrado. Fuente: job.M.CplexUsage.detailed.csv, columna BudgetUsed (acumulado) de la última llamada por instancia; "
              "mismo criterio y mismos valores que la Tabla tab:ganancia-tiempo del Cap. 4. "
              "Baseline = CPLEX puro por instancia (out/baseline/cplex_baseline_results.csv, reloj de pared).",
              lambda r, inst, i: round(r["t"].get(inst, 0.0), 3), 60)

# Sheet: tiempo total del algoritmo (wall clock)
ws4 = wb.create_sheet("Tiempo total algoritmo")
summary_sheet(ws4,
              "Tiempo TOTAL de ejecución del mejor algoritmo de cada grupo sobre cada instancia (s, reloj de pared): desde que el árbol "
              "empieza a evaluarse en la instancia hasta que termina, incluyendo heurísticas, CPLEX y todo el árbol. "
              "Fuente: job.M.MISPResults.out, columna 3 (ms), cronómetro System.nanoTime() en PDPProblemEvo.java:191-201. "
              "Baseline = CPLEX puro por instancia, también reloj de pared. Medido con 6 individuos evaluándose en paralelo en la misma máquina.",
              lambda r, inst, i: round(r["wall"][i] / 1000, 3), 75)

# Sheet: identificación del mejor algoritmo por grupo
ws2 = wb.create_sheet("Mejor algoritmo por grupo")
hdr2 = ["Grupo", "Condición", "Ejecución (1-5)", "Carpeta", "Gen (hoja, 1-100)", "Gen (archivo, 0-99)", "BestERP", "BestHits", "Tamaño árbol",
        "ID individuo (ECJ)", "Llamadas CPLEX (8 inst.)", "Tiempo CPLEX total (s)", "Tiempo evaluación completa (s, pared)"]
for c, h in enumerate(hdr2, 1):
    ws2.cell(1, c).value = h
style_header(ws2, 1, len(hdr2))
for i, r in enumerate(rows, 2):
    vals = [f"grupo{r['g']}", r["label"], r["ejec"], f"out/results/grupo{r['g']}/evolution{r['evo']}", r["gen"], r["gen"] - 1,
            r["erp"], r["hits"], r["size"], r["ind"], sum(r["calls"].values()), round(sum(r["t"].values()), 3), round(sum(r["wall"]) / 1000, 3)]
    for c, v in enumerate(vals, 1):
        ws2.cell(i, c).value = v
        ws2.cell(i, c).border = border
    ws2.cell(i, 7).number_format = "0.000000"
autowidth(ws2)
ws2.cell(len(rows) + 3, 1).value = ("Nota: la hoja 'Estadísticas Promedio y Mejor' numera generaciones 1-100; los archivos MISPResults.out y "
                                    "CplexUsage.detailed.csv numeran 0-99. Gen archivo = Gen hoja - 1.")

# Sheet: detalle por instancia
ws3 = wb.create_sheet("Detalle por instancia")
hdr3 = ["Grupo", "Condición", "Instancia", "Óptimo", "Costo obtenido", "Error relativo", "Llamadas CPLEX", "Tiempo CPLEX (s)",
        "Suma bruta TimeUsed (s)", "Secuencias registradas", "Tiempo evaluación completa (s, pared)", "Baseline CPLEX (s)"]
for c, h in enumerate(hdr3, 1):
    ws3.cell(1, c).value = h
style_header(ws3, 1, len(hdr3))
rr = 2
for r in rows:
    for k, inst in enumerate(INSTANCES):
        cost, opt = r["prof"][k]
        vals = [f"grupo{r['g']}", r["label"], inst.replace(".txt", ""), opt, cost, r["rel"][k], r["calls"].get(inst, 0),
                round(r["t"].get(inst, 0.0), 3), round(r["raw"].get(inst, 0.0), 3), r["seqs"].get(inst, 0),
                round(r["wall"][k] / 1000, 3), round(bl[inst], 3)]
        for c, v in enumerate(vals, 1):
            ws3.cell(rr, c).value = v
            ws3.cell(rr, c).border = border
        ws3.cell(rr, 6).number_format = "0.000000"
        rr += 1
autowidth(ws3)
ws3.cell(rr + 1, 1).value = ("Nota: 'Tiempo CPLEX (s)' usa el BudgetUsed acumulado de la última secuencia de llamadas por instancia (criterio de la tesis). "
                             "'Suma bruta TimeUsed' suma todas las filas atribuidas al individuo; difiere solo cuando 'Secuencias registradas' > 1, "
                             "lo que ocurre porque el logger es un singleton compartido por los 6 hilos de evaluación (evalthreads=6) y puede atribuir "
                             "llamadas de otro individuo concurrente. Tiempos CPLEX en CPU (ClockType=1, Threads=2); tiempo de evaluación y baseline en reloj de pared.")

# Sheet: Léeme (reading guide, first position)
ws0 = wb.create_sheet("Léeme", 0)
ws0.column_dimensions["A"].width = 34
ws0.column_dimensions["B"].width = 110


def sec(title):
    r = ws0.max_row + 2
    ws0.cell(r, 1).value = title
    ws0.cell(r, 1).font = Font(bold=True, size=12)
    ws0.cell(r, 1).fill = hdr_fill
    ws0.cell(r, 2).fill = hdr_fill


def kv(k, v):
    r = ws0.max_row + 1
    ws0.cell(r, 1).value = k
    ws0.cell(r, 1).font = bold
    ws0.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws0.cell(r, 2).value = v
    ws0.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)


ws0.cell(1, 1).value = "Tiempo de ejecución del mejor algoritmo generado por cada condición experimental"
ws0.cell(1, 1).font = Font(bold=True, size=14)
ws0.cell(2, 1).value = ("Tesis de Magíster, Matías Yáñez (USACH). Generación automática de algoritmos (GP) para el VRPSPD con un terminal "
                        "exacto (CPLEX) de presupuesto acotado. Datos extraídos directamente de los archivos de salida del experimento.")
ws0.cell(2, 1).alignment = Alignment(wrap_text=True)
ws0.merge_cells("A2:B2")
ws0.row_dimensions[2].height = 35

sec("1. Qué contiene este archivo")
kv("Pregunta que responde", "¿Cuánto tiempo tarda el MEJOR algoritmo generado en cada condición experimental en resolver cada una de las 8 instancias, "
   "y cuánto de ese tiempo corresponde a CPLEX? Se compara contra el tiempo que tarda CPLEX puro en resolver la misma instancia de forma exacta (baseline).")
kv("Condiciones experimentales", "Seis grupos, definidos por el presupuesto de tiempo que el terminal CPLEX puede usar en cada instancia, como porcentaje del tiempo baseline de esa instancia: "
   "B0 = 0 % (sin CPLEX, no aparece aquí porque su tiempo de CPLEX es cero), B10 = 10 %, B25 = 25 %, B50 = 50 %, B75 = 75 %, B100 = 100 %. "
   "En las carpetas del experimento: grupo1 = B10, grupo2 = B25, grupo3 = B50, grupo4 = B75, grupo5 = B100.")
kv("Qué es 'el mejor algoritmo'", "Cada condición se ejecutó 5 veces (5 evoluciones independientes de 100 generaciones y 50 individuos). El mejor algoritmo de una condición es el "
   "individuo con menor BestERP (error relativo promedio sobre las 8 instancias) observado en CUALQUIER generación de CUALQUIERA de las 5 ejecuciones. "
   "Es el mismo individuo que reporta la Tabla 'Mejor individuo observado por condición' del Capítulo 4.")
kv("Instancias", "Las 8 instancias del protocolo (conjunto 3C, 20 clientes, capacidades 50/66/80). Son las mismas usadas para evolucionar y para medir. "
   "El baseline es el tiempo de CPLEX puro resolviendo cada instancia hasta el óptimo.")

sec("2. Las dos medidas de tiempo (importante)")
kv("Tiempo total del algoritmo (hoja 'Tiempo total algoritmo')",
   "Tiempo de reloj de pared desde que el algoritmo generado empieza a ejecutarse sobre una instancia hasta que termina. Incluye TODO el árbol: heurísticas constructivas, "
   "búsqueda local, llamadas a CPLEX y reparación. Es la medida directa de 'cuánto tarda el algoritmo'. "
   "Se mide con System.nanoTime() dentro del propio hilo de evaluación, por lo que su atribución al individuo es confiable. "
   "Fuente: job.M.MISPResults.out, columna 3 (milisegundos).")
kv("Tiempo de CPLEX (hoja 'Tiempo CPLEX por instancia')",
   "Solo el tiempo consumido por el solver CPLEX dentro del algoritmo, sumando todas las llamadas parciales que el árbol hizo al evaluar la instancia. "
   "Se mide en tiempo de CPU (parámetro ClockType=1 de CPLEX, con Threads=2), NO en reloj de pared. Es la medida que usa la Tabla de ganancia en tiempo del Capítulo 4, "
   "y reproduce exactamente sus valores. Fuente: job.M.CplexUsage.detailed.csv.")
kv("Por qué el tiempo de CPLEX puede superar al tiempo total",
   "Porque están en relojes distintos. El tiempo de CPLEX es CPU acumulado con 2 hilos de CPLEX, en un proceso donde además se evaluaban 6 individuos en paralelo; "
   "el tiempo total es reloj de pared. En una misma instancia el tiempo de CPLEX (CPU) puede ser mayor que el tiempo total (pared). "
   "Por eso la comparación 'cuánto tarda el algoritmo' debe leerse en la hoja 'Tiempo total algoritmo'.")
kv("Condición de medición", "Los tiempos se midieron con 6 individuos evaluándose simultáneamente en la misma máquina (evalthreads=6). Con contención de CPU, "
   "el tiempo total medido es una cota superior del tiempo que tardaría el algoritmo ejecutándose solo. La ganancia frente al baseline es, por tanto, conservadora.")

sec("3. Guía hoja por hoja")
kv("Tiempo CPLEX por instancia", "Tabla resumen: filas = 8 instancias; columnas = baseline y las 5 condiciones híbridas. Cada celda es el tiempo de CPLEX (s) que usó el mejor algoritmo de esa "
   "condición en esa instancia. Al pie: total, ganancia en segundos (baseline menos total) y ganancia porcentual. Coincide con la Tabla de ganancia del Capítulo 4.")
kv("Tiempo total algoritmo", "Misma estructura que la anterior, pero con el tiempo total de ejecución del algoritmo (reloj de pared, s). Esta es la hoja que responde "
   "'cuánto tarda el mejor algoritmo en resolver cada instancia'.")
kv("Mejor algoritmo por grupo", "Una fila por condición. Identifica al mejor algoritmo (en qué ejecución y generación apareció, su identificador interno en ECJ) y sus totales sobre las 8 instancias: "
   "BestERP, BestHits (instancias en que igualó la solución de referencia), tamaño del árbol, número de llamadas a CPLEX, tiempo de CPLEX total y tiempo total de ejecución.")
kv("Detalle por instancia", "Una fila por cada par (condición, instancia): 5 x 8 = 40 filas. Es la tabla de origen de las dos hojas resumen. Incluye la calidad de la solución "
   "(óptimo, costo obtenido, error relativo), el número de llamadas a CPLEX, ambos tiempos y el baseline.")

sec("4. Glosario de columnas")
kv("Baseline CPLEX (s)", "Tiempo de reloj de pared que tardó CPLEX puro en resolver la instancia hasta el óptimo (out/baseline/cplex_baseline_results.csv). Referencia de comparación.")
kv("Ejecución (1-5)", "Cuál de las 5 evoluciones independientes de la condición contiene al mejor algoritmo. Numeración 1-5 como en la tesis; la carpeta correspondiente es evolution0-4.")
kv("Gen (hoja) / Gen (archivo)", "Generación en que apareció el mejor algoritmo. Las hojas 'Estadísticas Promedio y Mejor' numeran 1-100; los archivos crudos numeran 0-99. Gen archivo = Gen hoja - 1.")
kv("BestERP", "Error relativo promedio del algoritmo sobre las 8 instancias: promedio de |costo obtenido - óptimo| / óptimo. 0 = óptimo en todas.")
kv("BestHits", "Número de instancias (de 8) en que el algoritmo alcanzó exactamente la solución de referencia (error relativo = 0).")
kv("Tamaño árbol", "Número de nodos del árbol sintáctico del algoritmo generado.")
kv("ID individuo (ECJ)", "Identificador interno del individuo en el framework ECJ. Sirve para ubicarlo en los archivos crudos; no tiene otro significado.")
kv("Óptimo / Costo obtenido / Error relativo", "Óptimo = costo de la solución de referencia de la instancia. Costo obtenido = costo de la solución que construyó el algoritmo. Error relativo = (costo - óptimo) / óptimo.")
kv("Llamadas CPLEX", "Cuántas veces el algoritmo invocó al terminal CPLEX al resolver esa instancia.")
kv("Tiempo CPLEX (s)", "Tiempo de CPU consumido por CPLEX en esa instancia (suma de las llamadas). Criterio: valor acumulado (BudgetUsed) de la última llamada registrada, igual que en la tesis.")
kv("Suma bruta TimeUsed (s) / Secuencias registradas", "Columnas de control. Suma directa de todas las filas atribuidas al individuo en el registro de CPLEX y cuántas secuencias de llamadas aparecen. "
   "Normalmente 'Secuencias' = 1 y ambos tiempos coinciden. Cuando 'Secuencias' > 1 (ocurre en 3 instancias de B10) el registro mezcló llamadas de otro individuo que se evaluaba en paralelo, "
   "y se conserva el criterio de la tesis. Ver nota al pie de la hoja.")
kv("Tiempo evaluación completa (s, pared) / Tiempo total", "Tiempo total de ejecución del algoritmo sobre la instancia, de principio a fin, en reloj de pared. Ver sección 2.")
kv("Ganancia vs baseline (s) / Ganancia (%)", "Cuánto tiempo se ahorra frente a resolver la instancia con CPLEX puro: baseline menos total, y esa diferencia dividida por el baseline.")

sec("5. Fuentes")
kv("Archivos de origen", "ayud.-Investigacion/out/results/grupoN/evolutionM/: job.M.EstadisticaProm&Mej.csv (identificación del mejor), job.M.MISPResults.out (calidad y tiempo total por instancia), "
   "job.M.CplexUsage.detailed.csv (tiempo de CPLEX). ayud.-Investigacion/out/baseline/cplex_baseline_results.csv (baseline).")
kv("Reproducibilidad", "El archivo se genera con un script a partir de esos archivos, sin edición manual. Los valores de la hoja 'Tiempo CPLEX por instancia' fueron verificados celda a celda contra la Tabla de ganancia del Capítulo 4.")
for r in range(3, ws0.max_row + 1):
    if ws0.cell(r, 2).value:
        ws0.row_dimensions[r].height = max(30, 15 * (len(str(ws0.cell(r, 2).value)) // 105 + 1))

wb.save(OUT)
print("OK ->", OUT)
